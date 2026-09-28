# -*- coding: utf-8 -*-
"""Generate Confluence-importable Word (.docx) for Java Upgrade Agent."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Java-Upgrade-Agent-Confluence-Dokumani.docx"

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

for i in range(1, 4):
    hs = doc.styles[f"Heading {i}"]
    hs.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    hs.font.name = "Calibri"


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def p(t):
    doc.add_paragraph(t)


def bullet(t):
    doc.add_paragraph(t, style="List Bullet")


def num(t):
    doc.add_paragraph(t, style="List Number")


def code(t):
    para = doc.add_paragraph()
    run = para.add_run(t)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


title = doc.add_heading("Java Upgrade Agent", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Ürün ve Teknik Wiki Dökümanı")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
ir = meta.add_run("Desktop markası: Paytion Java Upgrade Desktop Agent")
ir.italic = True

p("Repo: https://github.com/dgknbingol/java-upgrade-agent")
p(
    "Kapsam: Maven tabanlı Java projelerini GitHub Copilot CLI ile hedef JDK sürümüne "
    "(varsayılan 21) otomatik yükseltme."
)
p(
    "Not: Bu Word dosyasını Confluence Space içinde Import → Word Document ile doğrudan "
    "içe aktarabilirsiniz. Heading stilleri Confluence başlık yapısına dönüşür; tablolar korunur."
)

h1("1. Özet (Executive Summary)")
p(
    "Java Upgrade Agent, şirket içi Java uygulama portföyünde manuel, yavaş ve hataya açık "
    "JDK yükseltme işini yapay zekâ destekli, doğrulanmış bir otomasyon pipeline’ına çeviren "
    "bir araçtır."
)
p('Klasik “pom.xml’de java.version değiştir” yaklaşımının aksine agent:')
bullet(
    "Kaynak kod, test, bağımlılık/BOM, plugin, Docker/CI ve dokümantasyonu kapsayan "
    "üretim kalitesinde migration hedefler"
)
bullet("Migration öncesi uyumluluk analizi yapar")
bullet("Değişiklikleri fazlı Copilot prompt’larıyla uygular")
bullet(
    "mvn clean install ile build doğrulaması + başarısızlıkta otomatik düzeltme döngüsü çalıştırır"
)
bullet("Migration borcunu katalog tabanlı validasyon ile kontrol eder")
bullet("(Desktop) Spring Boot uygulamayı ayağa kaldırarak smoke / runtime doğrulaması yapar")
bullet(
    "Sonucu rapor + git diff ile geliştiriciye sunar; branch push manueldir, otomatik PR açılmaz"
)
p("Desteklenen proje tipi: Yalnızca Maven. Gradle desteklenmez.")

h1("2. Nedir?")
p(
    "Java Upgrade Agent, üç parçalı bir monorepo içinde sunulan AI-assisted Java migration "
    "platformudur:"
)
add_table(
    ["Bileşen", "Açıklama"],
    [
        ["frontend/", "Merkezi web UI (React + Vite) — OpenShift’te tüm ekiplerin açtığı arayüz"],
        ["backend/", "Geliştirici laptop’unda çalışan lokal agent API (Express, port 4000)"],
        [
            "desktop-agent/",
            "Bağımsız Windows Electron uygulaması — UI + pipeline tek pakette; daha zengin özellik seti",
        ],
    ],
)
p("İki kullanım modeli vardır:")
num("Merkezi UI + Lokal Agent — UI OpenShift’te; git/mvn/copilot geliştirici makinesinde")
num("Desktop Agent — Tek .exe / Electron penceresi; harici backend yok")
p(
    "Ortak fikir: Kod, araçlar ve workspace şirket laptop’unda kalsın; merkezi sunucuya "
    "git/mvn/copilot kurma ve workspace biriktirme ihtiyacı olmasın."
)

h1("3. Ne işe yarar? Hangi problemi çözer?")
h2("3.1 Manuel yükseltmenin maliyeti")
p(
    "Kurumsal Java ekiplerinde JDK yükseltmesi (özellikle 8/11/17 → 17/21/25) tipik olarak "
    "şunları gerektirir:"
)
bullet("Compiler / release ayarları")
bullet("Spring Boot 2 → 3, Security 5 → 6")
bullet("javax.* → jakarta.*")
bullet("JUnit 4 → 5, PowerMock kaldırma")
bullet("HttpClient 4 → 5, Log4j1 → Log4j2/SLF4J")
bullet("Hibernate, Jackson, Swagger/Springfox, JDBC driver güncellemeleri")
bullet("Kaldırılmış/deprecated JDK API’leri (SecurityManager, finalize, Unsafe, …)")
bullet("Docker image tag’leri, CI java-version, README")
p(
    "Bu iş kıdemli Java bilgisi ister; çok dosyalı, çok modüllü repolarda kolay kaçırılır; "
    "build kırılınca tekrarlayan düzelt–build–düzelt döngüsüne girer; PR review öncesi "
    "tutarlı rapor üretmek zordur."
)

h2("3.2 Agent’ın değer önerisi")
add_table(
    ["Manuel süreç", "Agent ile"],
    [
        ["Sürüm tespiti elle", "Analiz: pom (+ backend’de Dockerfile/CI)"],
        ["Risk listesi deneyime bağlı", "Katalog tabanlı MIGRATION_ANALYSIS.md"],
        ["Değişiklikler dağınık commit’ler", "Fazlı Copilot migration + upgrade branch"],
        ["Build hataları elle", "Maven + Copilot fix loop"],
        ["“Bitti mi?” belirsiz", "MIGRATION_VALIDATION.md"],
        ["Runtime sürprizi", "Desktop: Run App / smoke"],
        ["Merkezi sunucuda clone", "Lokal workspace"],
    ],
)

h1("4. Ne amaçla kullanılır? (Use cases)")
num("LTS JDK yükseltmesi — Hedef: 8, 11, 17, 21, 25 (UI LTS listesi; varsayılan 21)")
num(
    "Spring Boot 2 → 3 geçişinin JDK ile birlikte yürütülmesi — Özellikle Java 21 + Boot 2 "
    "uyumsuzluğu agent tarafından yakalanır"
)
num("Portföy modernizasyonu — Çok sayıda Maven servisini aynı standart pipeline ile yükseltme")
num("Ön keşif / impact analysis — Sadece Analyze ile mevcut Java/Maven/Spring Boot ve risk özeti")
num("Build-only doğrulama — Migration sonrası Run Build / Fix Build")
num("Runtime doğrulama (Desktop) — mvn geçtikten sonra uygulamanın ayağa kalkması")
num("PR hazırlığı — Upgrade branch push → insan PR açar")
p("Ne için değildir:")
bullet("Gradle projeleri")
bullet("Otomatik merge / otomatik production deploy")
bullet("Otomatik Pull Request oluşturma")
bullet("Genel amaçlı LLM chat veya kod üretim IDE’si")

h1("5. Dağıtım modelleri")
h2("5.1 Model A — Merkezi UI + Lokal Backend")
code(
    "[OpenShift Frontend]\n"
    "        │  HTTP (VITE_API_BASE = http://localhost:4000)\n"
    "        ▼\n"
    "[Geliştirici laptop’u — backend :4000]\n"
    "        │\n"
    "   git clone / checkout / branch\n"
    "   copilot CLI (migration + fix)\n"
    "   mvn clean install\n"
    "   git push (manuel tetikleme)\n"
    "        │\n"
    "   backend/workspaces/{jobId}/"
)
p("Avantajlar:")
bullet("Tek merkezi UI URL’si")
bullet("OpenShift pod’una Git/Maven/Copilot kurmaya gerek yok")
bullet("Workspace ve credential’lar laptop’ta kalır")
bullet("Prompt/config Config Server ile merkezileştirilebilir")
p("Dikkat:")
bullet("HTTPS UI → HTTP localhost çağrısında tarayıcı kısıtı olabilir")
bullet("CORS: server.cors.origin merkezi UI origin’ini içermelidir")
bullet("Her geliştirici lokalde backend çalıştırmalıdır")

h2("5.2 Model B — Desktop Agent (Electron)")
code(
    "[React Renderer] ──IPC──► [Electron Main / jobRunner]\n"
    "                              │\n"
    "                    git / mvn / gh / copilot\n"
    "                              │\n"
    "               workspaces/{jobId}/  veya seçilen klasör\n"
    "               .java-upgrade/ artifact’ları"
)
p("Avantajlar:")
bullet("Tek uygulama, backend ayrı kurulum yok")
bullet("Smoke / Run App, Rollback, Stop Job, Copilot model seçimi, pipeline ayarları UI")
bullet("Windows NSIS installer / portable .exe")
p("Kısıtlar (POC):")
bullet("Windows odaklı")
bullet(
    "Job state bellekte; uygulama restart’ında aktif job handle kaybolur "
    "(workspace diskte kalır)"
)
bullet("OpenShift / DB / Docker orchestration yok")

h1("6. Kullanılan teknolojiler")
h2("6.1 Yazılım stack")
add_table(
    ["Katman", "Teknoloji", "Versiyon / not"],
    [
        ["UI (web)", "React + Vite + TypeScript", "React 19, Vite 6, TS 5.8"],
        ["API (web agent)", "Node.js + Express + TypeScript", "Express 4, uuid, cors"],
        ["Desktop UI", "React + Vite", "Port 5174 (dev)"],
        ["Desktop runtime", "Electron + electron-builder", "Electron 34, builder 25"],
        ["AI motoru", "GitHub Copilot CLI (@github/copilot)", "copilot -p ..."],
        ["Build doğrulama", "Apache Maven", "Host PATH"],
        ["VCS", "Git", "Host PATH"],
        ["(Desktop) GitHub CLI", "gh", "Health check"],
        [
            "Konfigürasyon",
            "application.properties (+ Config Server)",
            "Spring Boot tarzı anahtarlar",
        ],
        ["Canlı log (web)", "Server-Sent Events (SSE)", "/api/jobs/:id/events"],
        ["Canlı log (desktop)", "Electron IPC events", "job:log, job:status"],
    ],
)

h2("6.2 Host bağımlılıkları (paketlenmez)")
add_table(
    ["Araç", "Web backend", "Desktop"],
    [
        ["Git", "Zorunlu", "Zorunlu"],
        ["Maven", "Zorunlu", "Zorunlu"],
        ["JDK (≥ hedef major)", "Zorunlu", "Zorunlu (hedefe göre validate)"],
        ["@github/copilot", "Zorunlu (Node 22+)", "Zorunlu"],
        ["gh", "README’de zorunlu değil", "Health check’te var"],
        ["Copilot aboneliği + login", "Zorunlu", "Zorunlu"],
    ],
)

h2("6.3 Kullanılmayanlar")
bullet("Doğrudan OpenAI / Anthropic SDK")
bullet("Veritabanı")
bullet("Gradle")
bullet("Otomatik PR API’si (GitHub/GitLab)")

h1("7. Mimari")
h2("7.1 Yüksek seviye")
code(
    "UI: OpenShift Frontend  VEYA  Electron Renderer\n"
    "                 │ HTTP :4000  veya  IPC\n"
    "                 ▼\n"
    "Job Runner\n"
    "Analyze → [Turlar: Analyze → Copilot×4 → Maven+fix → Valid]\n"
    "(+ Desktop: Run App / Stop / Rollback)\n"
    "       │\n"
    "  Git / Copilot CLI / Maven / Smoke (desktop)\n"
    "       │\n"
    "Workspace + rapor/diff artifact’ları\n"
    "Manuel Push → insan PR açar"
)

h2("7.2 Backend (web agent) modülleri")
add_table(
    ["Modül", "Görev"],
    [
        ["server.ts", "REST + SSE"],
        ["jobRunner.ts", "Upgrade / build / push orkestrasyonu"],
        ["jobStore.ts", "Bellek içi job state"],
        ["repoAnalyzer / javaVersionAnalyzer", "Java sürüm tespiti (pom, Dockerfile, GHA)"],
        ["migrationAnalyzer", "Uyumluluk taraması"],
        ["migrationCatalog", "Pattern / kategori kütüphanesi"],
        ["migrationValidator", "Post-build doğrulama"],
        ["pomUpgrader", "Deterministik POM baseline"],
        ["migrationPrompts / javaUpgradePrompt", "Prompt üretimi"],
        ["artifacts.ts", "Diff + fallback rapor"],
        ["prerequisites.ts", "Araç kontrolü"],
        ["config/*", "Properties + Config Server"],
    ],
)

h2("7.3 Desktop servisleri")
p("Web ile aynı çekirdek + ekler:")
add_table(
    ["Servis", "Görev"],
    [
        ["gitService", "clone, branch, diff, commit+push, rollback"],
        ["mavenService", "analyzePom, clean install, package"],
        ["copilotService", "Copilot spawn + task MD"],
        ["copilotModelService", "Model listesi / seçimi"],
        ["healthService", "Önkoşul kontrolü"],
        ["smokeService", "Spring Boot startup smoke"],
        ["projectLocalProperties", "Lokal properties enjekte"],
        ["commandRunner / spawnUtil", "Process yönetimi"],
        ["jobCancellation", "Job stop"],
    ],
)

h1("8. Uçtan uca iş akışı")
h2("8.1 Kullanıcı adımları")
num("Health / Prerequisites — Git, Java, Maven, Copilot (desktop: gh)")
num("Repo URL + source branch (veya lokal klasör)")
num("Analyze — mevcut Java / Maven / Spring Boot")
num("Hedef Java seçimi (varsayılan 21)")
num("Start Upgrade")
num("Log / rapor / diff inceleme")
num("İsteğe bağlı: Run Build, Run App (desktop)")
num("Push Branch")
num("İsteğe bağlı: Rollback (desktop)")
num("Git hosting üzerinde manuel PR")

h2("8.2 Start Upgrade pipeline (çekirdek)")
code(
    "Workspace hazırla (clone veya local)\n"
    "  → feature/java-{version}-upgrade branch (çakışmada -2, -3…)\n"
    "  → Baseline analyze → MIGRATION_ANALYSIS.md\n"
    "  → for round = 1..maxMigrationRounds (varsayılan 3):\n"
    "        yeniden analyze\n"
    "        round 1:\n"
    "          pomUpgrader (baseline)\n"
    "          Copilot faz 1: build-system\n"
    "          Copilot faz 2: dependencies-ecosystem\n"
    "          Copilot faz 3: source-code\n"
    "          Copilot faz 4: infrastructure (+ rapor)\n"
    "        round 2+: continuation prompt (validation + build log)\n"
    "        değişiklik yoksa ve upgrade gerekiyorsa fail\n"
    "        mvn clean install\n"
    "          başarısızsa Copilot build-fix (maxBuildFixAttempts, varsayılan 2)\n"
    "        build OK → validateMigration → MIGRATION_VALIDATION.md\n"
    "        validation pass → bitir\n"
    "        değilse soft-pass kuralları veya sonraki tur\n"
    "  → artifact sync (report + git diff)"
)

h2("8.3 Dört Copilot migration fazı")
add_table(
    ["Faz", "Odak"],
    [
        ["build-system", "POM’lar, compiler, plugin’ler, enforcer"],
        ["dependencies-ecosystem", "BOM, framework, javax→jakarta, legacy deps"],
        ["source-code", "main/test/config, JDK API, JUnit 4→5"],
        ["infrastructure", "Docker, CI, script, docs + migration report"],
    ],
)

h2("8.4 Desktop ek akışlar")
bullet(
    "Run App / Smoke: Spring Boot modül tespiti → jar / maven / auto → başarı "
    "pattern’leri (Started … in) → hata olursa runtime fix + rebuild"
)
bullet("Stop: Çalışan process’leri iptal")
bullet("Rollback: Migration değişikliklerini geri al / source branch’e dön")

h1("9. Sağladığı olanaklar (Capabilities)")
num("Otomatik sürüm tespiti — Java, Maven compiler, Spring Boot parent")
num("Katalog tabanlı risk analizi — import/API/pom pattern’leri")
num("Fazlı AI migration — tek prompt yerine 4 aşamalı kontrollü değişiklik")
num("Çok turlu dış döngü — analyze → migrate → build → validate (max 3)")
num("İç Maven fix loop — build kırılınca Copilot ile onarım")
num("Deterministik pomUpgrader — Copilot eksik bırakırsa Java property/plugin baseline")
num("Validation gate — POM mismatch, kalan legacy namespace, Boot 2 + Java 21 vb.")
num("Canlı log streaming — SSE veya IPC")
num("Görsel diff özeti + raw diff")
num("Migration raporları — analysis / validation / report markdown")
num("Manuel push + güvenli branch isimlendirme")
num("Config Server ile merkezi prompt/ayar (web backend)")
num("Desktop: smoke, rollback, stop, model seçimi, pipeline settings UI, lokal folder pick")
num("Workspace izolasyonu — job başına klasör")

h1("10. Migration katalog kapsamı")
p("Agent’ın taradığı kategori örnekleri:")
add_table(
    ["Kategori", "Örnek"],
    [
        ["legacy-namespace", "javax.servlet, javax.persistence, javax.validation, … → jakarta.*"],
        ["removed-jdk-api", "SecurityManager, Thread.stop, …"],
        ["deprecated-jdk-api", "finalize(), Class.newInstance(), legacy URL"],
        ["internal-jdk-api", "sun.misc.Unsafe, com.sun.*"],
        ["legacy-test-api", "JUnit 4, PowerMock, eski Mockito runners"],
        ["legacy-logging", "Log4j1, commons-logging"],
        ["legacy-http-client", "Apache HttpClient 4"],
        ["incompatible-dependency", "Jackson 1 (codehaus), Springfox, eski Swagger annotations"],
        [
            "spring-boot-legacy",
            "WebSecurityConfigurerAdapter, antMatchers, eski actuator property’ler",
        ],
        ["reflection-instrumentation", "setAccessible(true) / module sorunları"],
        ["legacy-serialization", "ObjectInputStream, XStream"],
    ],
)
p(
    "Tasarım ilkesi (prompt’larda): En küçük güvenli yol; hedef JDK için zorunlu olmayan "
    "framework yükseltmelerini yapma."
)

h1("11. AI / Copilot entegrasyonu")
bullet("Araç: npm install -g @github/copilot → copilot binary")
bullet("Kimlik: copilot login + aktif Copilot aboneliği")
bullet(
    "Tipik bayraklar: -p, --model, --allow-all-tools, --allow-all-paths, "
    "--no-ask-user, --autopilot, --stream on"
)
bullet(
    "Desktop pattern: Uzun talimat .java-upgrade/JAVA_UPGRADE_TASK.md "
    "(veya BUILD_FIX_TASK.md / RUNTIME_FIX_TASK.md); CLI kısa -p ile dosyayı açtırır"
)
bullet(
    "Backend: Prompt -p ile / dosyaya yazma; şablonlar application.properties içinde "
    "de tutulabilir"
)
bullet("Placeholder’lar: {sourceJavaVersion}, {targetJavaVersion}, {buildLog}, …")
bullet("Model: Desktop’ta copilot.model (varsayılan auto) + model listesi IPC")
p(
    "Not: Eski architecture.md metnindeki “gh copilot suggest” ifadesi güncel kodla "
    "uyumsuz olabilir; migration @github/copilot CLI ile yapılır."
)

h1("12. Artifact’lar ve çıktılar")
add_table(
    ["Artifact", "Tipik konum (desktop)", "İçerik"],
    [
        ["MIGRATION_REPORT.md", ".java-upgrade/", "Kararlar, dosyalar, riskler, blocker’lar"],
        ["MIGRATION_ANALYSIS.md", ".java-upgrade/", "Uyumluluk bulguları"],
        ["MIGRATION_VALIDATION.md", ".java-upgrade/", "Pass/fail"],
        ["Task MD dosyaları", ".java-upgrade/", "Copilot talimatları"],
        ["RUN_APP_LOG.txt", "smoke", "Startup log"],
        ["Git diff", "UI / API", "Source branch’e göre"],
        ["Workspace", "workspaces/{jobId}/", "Clone"],
    ],
)
p(r"Packaged desktop workspace: %APPDATA%\java-upgrade-desktop-agent\workspaces\{jobId}\ ")
p("Push commit mesajı (desktop): chore(java): upgrade to Java {version}")
p("PR: Kullanıcı GitHub/GitLab’da manuel açar.")

h1("13. API ve IPC özeti")
h2("13.1 Backend REST")
add_table(
    ["Method", "Endpoint", "Açıklama"],
    [
        ["GET", "/api/config", "Aktif config özeti"],
        ["GET", "/api/prerequisites", "Araç kontrolü"],
        ["POST", "/api/pick-folder", "Klasör seçimi"],
        ["POST", "/api/analyze", "Java sürüm analizi"],
        ["POST", "/api/jobs/start", "Upgrade job"],
        ["GET", "/api/jobs/:jobId/events", "SSE log/status"],
        ["GET", "/api/jobs/:jobId/report", "Migration report"],
        ["GET", "/api/jobs/:jobId/diff", "Git diff"],
        ["POST", "/api/jobs/:jobId/build", "Maven rebuild"],
        ["POST", "/api/jobs/:jobId/push", "Branch push"],
    ],
)
h2("13.2 Desktop IPC (seçilmiş)")
p(
    "health:check, job:analyze, job:start, job:runBuild, job:runStartup, job:push, "
    "job:rollback, job:stop, job:getArtifacts, job:getLiveReport, job:getLiveDiff, "
    "copilot:listModels, app:savePipelineConfig, …"
)

h1("14. Yapılandırma")
h2("14.1 Yükleme sırası (backend)")
num("config/application.properties")
num("config/application-local.properties (gitignore, kişisel)")
num("Spring Cloud Config Server")
num("Env override (PORT, SERVER_CORS_ORIGIN, …)")

h2("14.2 Önemli anahtarlar")
add_table(
    ["Anahtar", "Varsayılan", "Anlam"],
    [
        ["server.port", "4000", "Backend port"],
        ["server.cors.origin", "http://localhost:5173", "CORS"],
        ["job.max-build-fix-attempts", "2", "İç Maven fix turu"],
        ["job.max-migration-rounds", "3", "Dış migration turu"],
        ["job.maven-build-log-tail-chars", "6000", "Fix prompt’a giden log kuyruğu"],
        ["job.upgrade-branch-pattern", "feature/java-{version}-upgrade", "Branch şablonu"],
        ["copilot.command", "copilot", "Binary adı"],
        ["copilot.idle-heartbeat-seconds", "20", "Idle heartbeat"],
        ["prompt.java-upgrade", "(uzun şablon)", "Migration prompt"],
        ["prompt.maven-fix", "(şablon)", "Build fix prompt"],
    ],
)
p(
    "Desktop ek anahtarlar: job.smoke-run-enabled, job.smoke-run-timeout-seconds, "
    "job.max-smoke-fix-attempts, job.startup-run-mode (jar|maven|auto), copilot.model."
)

h1("15. Verimlilik ve ROI")
h2("15.1 Repodaki durum")
p(
    "Kod tabanında ölçülmüş KPI yok (saat tasarrufu, başarı oranı, PR cycle time "
    "dashboard’u vb. yok). Bu yüzden Confluence’da “%X arttı” iddiası yalnızca kendi "
    "ölçümünüzle yazılmalıdır."
)

h2("15.2 Nitel (documented / implied) faydalar")
add_table(
    ["Alan", "Etki"],
    [
        ["Keşif süresi", "Analyze + katalog → risk listesi dakikalar içinde"],
        ["Uygulama süresi", "4 faz + çok tur; gece/arka planda çalıştırılabilir"],
        ["Build onarım", "Otomatik fix loop → kıdemli müdahale ihtiyacını azaltır"],
        ["Kaçırılan pattern", "Katalog + validation gate"],
        ["Review kalitesi", "Report + diff ile PR hazırlığı"],
        ["Operasyon", "Lokal çalıştırma → merkezi disk/secret yükü yok"],
        ["Standartlaşma", "Aynı prompt/pipeline tüm ekiplerde"],
    ],
)

h2("15.3 Ölçüm şablonu (ekibin doldurması için)")
add_table(
    ["Metrik", "Manuel baseline", "Agent ile", "Not"],
    [
        ["Ortalama upgrade süresi (kişi-saat)", "örn. 3–10 gün", "ölçün", "Proje boyutuna göre"],
        ["İlk yeşil mvn clean install’a kadar tur", "manuel", "agent log", ""],
        ["PR review cycle", "", "", "Report/diff etkisi"],
        ["Kritik kaçırılan item (javax, Boot 2, …)", "", "", "Validation sonrası"],
        ["Başarı oranı (push’a kadar)", "", "", "Soft-pass ayrımı yapın"],
    ],
)
p(
    "Pratik anlatım (ölçüm yokken): Agent, Java yükseltmesini keşif → AI uygulama → "
    "build onarım → validasyon → (opsiyonel) smoke → manuel push zincirine indirger; "
    "tekrarlayan manuel işi ve kaçırılan uyumluluk risklerini azaltır. Sayısal ROI için "
    "ekip pilot ölçümü gerekir."
)

h1("16. Kurulum ve çalıştırma")
h2("16.1 Host kurulum (Windows örnek)")
add_table(
    ["Araç", "Kontrol", "Kurulum örneği"],
    [
        ["Git", "git --version", "git-scm.com"],
        ["JDK 21+", "java --version", "winget install EclipseAdoptium.Temurin.21.JDK"],
        ["Maven", "mvn --version", "choco install maven -y"],
        ["Copilot CLI", "copilot --version", "npm install -g @github/copilot (Node 22+)"],
        ["GitHub CLI (desktop)", "gh --version", "winget install GitHub.cli"],
    ],
)
p("Kurulum sonrası terminal/Cursor restart; copilot ile login.")

h2("16.2 Web stack")
code(
    "cd backend && npm install && npm run dev    # :4000\n"
    "cd frontend && npm install && npm run dev   # :5173"
)

h2("16.3 Desktop")
code(
    "cd desktop-agent && npm install && npm run dev\n"
    "# paketleme:\n"
    "npm run dist          # NSIS installer\n"
    "npm run dist:dir      # unpacked exe\n"
    "npm run dist:portable # portable"
)

h1("17. Güvenlik ve operasyonel notlar")
bullet(
    "Desktop: contextIsolation: true, nodeIntegration: false; process spawn sadece main process"
)
bullet("Copilot --allow-all-tools / --autopilot → güvenilen workspace varsayımı")
bullet("Push asla otomatik değil")
bullet("Job state bellekte; restart’ta aktif job kaybolur, workspace kalır")
bullet("Credential’lar geliştirici makinesinde (Git SSH/HTTPS, Copilot auth)")
bullet(
    "HTTPS merkezi UI → localhost HTTP için Private Network Access / IT politikası "
    "kontrol edilmeli"
)

h1("18. Bilinen kısıtlar")
num("Sadece Maven")
num("Job persistence yok (DB yok)")
num("Otomatik PR yok")
num("Gradle / non-Java monorepo senaryoları desteklenmez")
num(
    "Boş diff: kaynak zaten hedef JDK’deyse Copilot değişiklik yapmayabilir "
    "(build yine geçebilir)"
)
num(
    "Desktop Windows odaklı; architecture.md bazı yerlerde eski gh copilot dilini "
    "taşıyabilir"
)
num(
    "AI çıktısı her zaman doğru olmayabilir → validation + insan review zorunlu "
    "kabul edilmeli"
)

h1("19. Proje klasör yapısı")
code(
    "java-upgrade-agent/\n"
    "├── README.md                 # Web: merkezi UI + lokal agent\n"
    "├── frontend/                 # React UI (:5173)\n"
    "├── backend/                  # Express agent (:4000)\n"
    "│   ├── config/application.properties\n"
    "│   ├── src/                  # server, jobRunner, migration*, …\n"
    "│   └── workspaces/\n"
    "└── desktop-agent/            # Electron (Paytion …)\n"
    "    ├── README.md\n"
    "    ├── architecture.md\n"
    "    ├── electron/             # main, services, prompts\n"
    "    ├── src/                  # renderer UI\n"
    "    ├── config/\n"
    "    ├── workspaces/\n"
    "    └── release/              # installer çıktıları"
)

h1("20. Tipik senaryo")
num("Geliştirici UI’ı açar, repo URL + develop girer.")
num("Analyze: Java 11, Spring Boot 2.7 tespit eder.")
num("Hedef Java 21 seçilir.")
num("Start Upgrade: feature/java-21-upgrade açılır.")
num(
    "Katalog javax/Spring Security legacy bulur; Copilot 4 fazda POM + kaynak + CI günceller."
)
num("mvn clean install kırılır → Copilot fix → yeşil build.")
num("Validation: Boot 2 + Java 21 residual uyarısı varsa ek tur.")
num("(Desktop) Run App: Started MyServiceApplication.")
num("Diff/report review → Push Branch → PR açılır.")

h1("21. Confluence sayfa önerisi")
num("Overview (bu döküman)")
num("Deployment models")
num("User guide (ekip üyesi)")
num("Admin guide (Config Server, CORS, OpenShift frontend)")
num("Desktop Agent packaging")
num("Architecture & pipeline deep dive")
num("Migration catalog reference")
num("Prompt governance")
num("Security & compliance")
num("Metrics & pilot results (ölçümle doldurulacak)")
num("FAQ / troubleshooting")

h1("22. FAQ")
p("Gradle destekleniyor mu? Hayır.")
p("PR otomatik açılır mı? Hayır; sadece branch push.")
p(
    "Neden lokal backend? Git/Maven/Copilot ve workspace’in laptop’ta kalması; "
    "OpenShift’e araç yükü taşımamak."
)
p("Job kayboldu? Backend/app restart job state’i siler; workspaces/ diskte kalır.")
p(
    "Hedef 25 seçtim, JDK 21 var? Health check hedef major’ı karşılamalı; JDK "
    "yükseltilmeli."
)
p(
    "Copilot değişiklik yapmadı? Kaynak zaten hedef sürümde olabilir; gerçek "
    "migration için eski JDK’li repo kullanın."
)

h1("Ek A — Confluence’a nasıl import edilir?")
num(
    "Confluence space’inizde Create / Import (veya Space tools → Content Tools → Import) "
    "seçin."
)
num("Word Document (.docx) seçin.")
num("Bu dosyayı yükleyin.")
num(
    "Başlık stilleri (Heading 1/2/3) Confluence sayfa/başlık yapısına dönüşür; tablolar "
    "korunur."
)
num(
    "İsterseniz import sonrası tek sayfa yerine Heading 1’leri ayrı child page’lere "
    "bölebilirsiniz."
)

doc.add_paragraph()
footer = doc.add_paragraph()
fr = footer.add_run(
    "Belge: Java Upgrade Agent — Confluence import için hazırlanmıştır. "
    "Kaynak: proje README + architecture + kod tabanı."
)
fr.italic = True

doc.save(OUT)
print(f"OK: {OUT}")
print(f"Size: {OUT.stat().st_size} bytes")
